"""可直接运行：python -m biaoshu --help"""
from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

from .core import allocate, initialize, load_project, scaffold, section_path, write_json
from .quality_gate import inspect_with_evidence
from .extraction import extract_document
from .score_candidates import discover_score_candidates
from .ai_tender_analysis import analyze_tender
from .review_import import review_template, reviewed_import
from .review_transaction import recover as recover_review
from .outline_tree import (outline_template, outline_plan, lock_outline,
                           walk_outline, validate_outline_state)
from .section_writer import write_sections, approve_section
from .section_transaction import recover as recover_section
from .consistency import analyze_consistency
from .response_trace import create_trace, audit_trace
from .semantic_review import review_responses
from .safe_report import write_new_report, validate_new_report_target


def extract_source(source: Path) -> str:
    """Compatibility helper; CLI also saves a machine-readable evidence index."""
    return extract_document(source).markdown


def export_docx(root: Path, dest: Path, draft: bool, *,
                index_paths: list[Path] | None = None,
                source_dir: Path | None = None):
    """Export only after structural and source-locked evidence checks.

    Direct Python callers are subject to the same official-export gate as the CLI.
    A draft remains explicitly marked unverified and cannot be mistaken for approval.
    """
    report = inspect_with_evidence(root, strict=True, index_paths=index_paths,
                                   source_dir=source_dir, require_evidence=not draft)
    if not report["ok"] and not draft:
        raise ValueError("导出被严格审查阻止；先修复错误，或使用 --draft 输出明确标记的草稿")
    try:
        from docx import Document
        from docx.shared import Pt
    except ImportError as exc:
        raise RuntimeError('DOCX 导出需要 pip install -e ".[docs]"') from exc
    state = load_project(root)
    tree_errors = validate_outline_state(root, state["outline.json"])
    if tree_errors:
        raise ValueError("目录层级或锁定状态无效：" + "; ".join(tree_errors))
    document = Document()
    title = state["project.json"].get("project_name") or "未命名项目"
    document.add_heading(f"{title} — 技术投标文件" + ("（草稿，未完成校验）" if draft else ""), 0)
    for level, kind, item in walk_outline(state["outline.json"]):
        if kind == "group":
            document.add_heading(str(item["title"]), level=min(level, 4))
            continue
        sec = item
        document.add_heading(str(sec["title"]), level=min(level, 4))
        path = section_path(root, sec["id"])
        if not path.exists():
            document.add_paragraph(f"[缺少正文：{sec['id']}]")
            continue
        for raw in path.read_text(encoding="utf-8").splitlines():
            line = raw.strip()
            if not line:
                continue
            if line.startswith("#"):
                text = line.lstrip("#").strip()
                if text and text != sec["title"]:
                    document.add_heading(text, level=min(len(line) - len(line.lstrip("#")) + level, 4))
            elif re.match(r"^[-*]\s+", line):
                document.add_paragraph(re.sub(r"^[-*]\s+", "", line), style="List Bullet")
            elif re.match(r"^\d+[.)]\s+", line):
                document.add_paragraph(re.sub(r"^\d+[.)]\s+", "", line), style="List Number")
            else:
                document.add_paragraph(line)
    for style in ("Normal",):
        document.styles[style].font.size = Pt(11)
    dest.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(dest))
    return report


def main(argv=None):
    parser = argparse.ArgumentParser(prog="biaoshu", description="技术标 Skill 的确定性工具层")
    sub = parser.add_subparsers(dest="command", required=True)
    command_root = ("init", "check", "audit", "allocate", "scaffold", "export")
    for name in command_root:
        cmd = sub.add_parser(name)
        cmd.add_argument("project_dir", type=Path)
        if name in ("check", "audit", "export"):
            cmd.add_argument("--index", action="append", type=Path, default=[],
                             help="证据索引 .index.json，可重复提供多个")
            cmd.add_argument("--source-dir", type=Path,
                             help="原始招标文件所在目录（必须与 --index 同时提供）")
        if name == "allocate":
            cmd.add_argument("--total", type=int, required=True, help="全文 AI 正文目标字数")
        if name == "export":
            cmd.add_argument("--output", type=Path, required=True)
            cmd.add_argument("--draft", action="store_true")
        if name in ("check", "audit"):
            cmd.add_argument("--output", type=Path, help="保存校验报告 JSON")
    ext = sub.add_parser("extract")
    ext.add_argument("source", type=Path)
    ext.add_argument("--output", type=Path, required=True)
    score_cmd = sub.add_parser("score-candidates", help="提取待核准的技术评分项候选")
    score_cmd.add_argument("--index", action="append", type=Path, required=True,
                           help="招标文件证据索引，可多次指定")
    score_cmd.add_argument("--source-dir", type=Path, required=True,
                           help="原始招标文件目录，用于验证 SHA-256")
    score_cmd.add_argument("--output", type=Path, required=True,
                           help="保存候选 JSON，不修改 requirements.json")
    ai_cmd = sub.add_parser("analyze-ai", help="调用显式指定的本地 AI 模型分析已索引的招标原文")
    ai_cmd.add_argument("--index", action="append", type=Path, required=True,
                        help="证据索引，可重复传入")
    ai_cmd.add_argument("--source-dir", type=Path, required=True,
                        help="必须包含未修改的原始招标文件")
    ai_cmd.add_argument("--model", required=True,
                        help="明确指定已安装的模型名称")
    ai_cmd.add_argument("--endpoint", default="http://127.0.0.1:11434/v1",
                        help="OpenAI 兼容 base URL；默认 Ollama 本地服务")
    ai_cmd.add_argument("--allow-remote", action="store_true",
                        help="明确允许将招标资料发送到远程模型服务；请先获得资料授权")
    ai_cmd.add_argument("--api-key-env",
                        help="可选：从指定环境变量读取 API key；不写入命令参数、日志或仓库")
    ai_cmd.add_argument("--timeout", type=float, default=90,
                        help="模型请求超时秒数（1-600）")
    ai_cmd.add_argument("--output", type=Path, required=True,
                        help="写入待人工确认建议 JSON，永不覆盖正式 requirements.json")
    reviewer = sub.add_parser("review-init", help="为 AI/规则候选创建人工审批清单")
    reviewer.add_argument("--suggestions", required=True, type=Path)
    reviewer.add_argument("--output", required=True, type=Path)
    importer = sub.add_parser("review-import", help="预览或登记明确核准的技术需求与一一响应章节")
    importer.add_argument("project_dir", type=Path)
    importer.add_argument("--suggestions", required=True, type=Path)
    importer.add_argument("--review", required=True, type=Path)
    importer.add_argument("--index", required=True, action="append", type=Path)
    importer.add_argument("--source-dir", required=True, type=Path)
    importer.add_argument("--apply", action="store_true",
                          help="经人工逐项审查并填写 acknowledgement 后才实际写入")
    importer.add_argument("--output", type=Path, help="可选预览/导入报告 JSON")
    outline_init = sub.add_parser("outline-init", help="生成保持章节 ID 的层级目录候选，不覆盖已有文件")
    outline_init.add_argument("project_dir", type=Path)
    outline_init.add_argument("--output", type=Path, required=True)
    outline_cmd = sub.add_parser("outline-plan", help="预览/应用目录层级方案（不修改 S# 和正文）")
    outline_cmd.add_argument("project_dir", type=Path)
    outline_cmd.add_argument("--proposal", type=Path, required=True)
    outline_cmd.add_argument("--apply", action="store_true")
    outline_cmd.add_argument("--output", type=Path)
    outline_lock = sub.add_parser("outline-lock", help="显式锁定一级目录名称、顺序和章节归属")
    outline_lock.add_argument("project_dir", type=Path)
    outline_lock.add_argument("--confirm", action="store_true")
    writing = sub.add_parser("write-sections", help="按评分项逐章节 AI 生成未批准草稿；已有正文不覆盖")
    writing.add_argument("project_dir", type=Path)
    writing.add_argument("--section", action="append", help="指定 S# 章节；不指定时按目录顺序处理全部")
    writing.add_argument("--index", action="append", type=Path, required=True)
    writing.add_argument("--source-dir", type=Path, required=True)
    writing.add_argument("--model", required=True)
    writing.add_argument("--endpoint", default="http://127.0.0.1:11434/v1")
    writing.add_argument("--allow-remote", action="store_true")
    writing.add_argument("--api-key-env")
    writing.add_argument("--timeout", type=float, default=90)
    approving = sub.add_parser("approve-section", help="人工逐项审查后批准一个指定章节")
    approving.add_argument("project_dir", type=Path)
    approving.add_argument("--section", required=True)
    approving.add_argument("--acknowledgement", required=True,
                           help="必须填写真实人工复核确认句，不得由模型代替")
    consistency = sub.add_parser("consistency-audit", help="本地跨章节承诺一致性审查")
    consistency.add_argument("project_dir", type=Path)
    consistency.add_argument("--output", type=Path, help="保存结构化审查报告 JSON")
    trace_init = sub.add_parser("trace-init", help="建立需逐项人工填写的评分响应段落台账，不覆盖旧记录")
    trace_init.add_argument("project_dir", type=Path)
    trace_audit = sub.add_parser("trace-audit", help="只读审查评分要求到正文原句的人工确认记录")
    trace_audit.add_argument("project_dir", type=Path)
    trace_audit.add_argument("--output", type=Path, help="保存本地 JSON 审查报告")
    semantic = sub.add_parser("semantic-review", help="本地 AI 逐条审阅已定位技术响应并提出建议")
    semantic.add_argument("project_dir", type=Path)
    semantic.add_argument("--index", action="append", type=Path, required=True)
    semantic.add_argument("--source-dir", type=Path, required=True)
    semantic.add_argument("--model", required=True)
    semantic.add_argument("--endpoint", default="http://127.0.0.1:11434/v1")
    semantic.add_argument("--allow-remote", action="store_true",
                          help="确认招标资料已获授权后才允许远程模型")
    semantic.add_argument("--api-key-env")
    semantic.add_argument("--timeout", type=float, default=90)
    semantic.add_argument("--output", type=Path, required=True,
                          help="只生成独立待复核建议 JSON，不修改项目或审批台账")
    recovery = sub.add_parser("review-recover", help="中断的审核导入按事务备份恢复，绝不合并未知外部改动")
    recovery.add_argument("project_dir", type=Path)
    recovery.add_argument("--confirm", action="store_true")
    recover_chapter = sub.add_parser("section-recover", help="人工确认后恢复中断的章节写入；拒绝覆盖未知修改")
    recover_chapter.add_argument("project_dir", type=Path)
    recover_chapter.add_argument("--section", required=True)
    recover_chapter.add_argument("--confirm", action="store_true")
    recover_chapter.add_argument("--unlock-stale", action="store_true",
                                 help="只在确认原写作进程已停止时解除残留锁")
    args = parser.parse_args(argv)
    try:
        if args.command in ("semantic-review", "analyze-ai", "score-candidates",
                            "review-init", "outline-init"):
            # Fail before model calls, and use atomic exclusive publishing
            # after the model returns to close the check/write race.
            validate_new_report_target(args.output)
        if args.command == "init":
            initialize(args.project_dir)
            print(f"已初始化：{args.project_dir}")
        elif args.command == "extract":
            extracted = extract_document(args.source)
            args.output.parent.mkdir(parents=True, exist_ok=True)
            args.output.write_text("\n".join(["# 招标原文提取", "", f"来源：{args.source.name}", "", extracted.markdown]), encoding="utf-8")
            index_path = args.output.with_suffix(".index.json")
            write_json(index_path, extracted.index)
            print(f"已提取：{args.output}；证据索引：{index_path}。请人工核对缺失文本和提取警告。")
        elif args.command == "semantic-review":
            result = review_responses(
                args.project_dir, args.index, args.source_dir,
                model=args.model, endpoint=args.endpoint,
                allow_remote=args.allow_remote, api_key_env=args.api_key_env,
                timeout=args.timeout)
            write_new_report(args.output, result)
            print(f"已生成 AI 逐项语义审稿建议 {result['statistics']['examined_links']} 项；"
                  f"潜在问题 {result['statistics']['potential_issues']} 项；"
                  f"结果 {args.output}。全部待人工复核，不影响正式批准。")
        elif args.command == "trace-init":
            result = create_trace(args.project_dir)
            print(f"已创建待人工填写的响应追溯台账：{result['path']}，"
                  f"需核对 {result['items']} 个需求与章节关联。")
        elif args.command == "trace-audit":
            result = audit_trace(args.project_dir)
            if args.output:
                write_json(args.output, result)
            print(f"评分响应定位审查：{result['status']}；"
                  f"已确认 {result['statistics']['confirmed_links']}/"
                  f"{result['statistics']['expected_links']} 项；"
                  f"待解决 {result['statistics']['errors']} 个。")
            for item in result["findings"]:
                print(f"{item['code']}：{item['message']} [{item['requirement_id']}/{item['section_id']}]")
            return 0 if result["ok"] else 1
        elif args.command == "consistency-audit":
            report = analyze_consistency(args.project_dir)
            if args.output:
                write_json(args.output, report)
            print(f"跨章节扫描：{report['statistics']['sections_inspected']} 节；"
                  f"明确承诺 {report['statistics']['commitments_detected']} 处；"
                  f"冲突 {report['statistics']['errors']} 个；"
                  f"提示 {report['statistics']['warnings']} 个。")
            for item in report["findings"]:
                print(f"[{item['severity']}] {item['code']}：{item['message']}")
            return 0 if report["ok"] else 1
        elif args.command == "write-sections":
            report = write_sections(args.project_dir, args.section, args.index,
                                    args.source_dir, model=args.model,
                                    endpoint=args.endpoint, allow_remote=args.allow_remote,
                                    api_key_env=args.api_key_env, timeout=args.timeout)
            for item in report["outcomes"]:
                print(f"章节 {item['section_id']}：{item['status']}")
            print(f"本轮处理 {report['completed']} 节。AI 新内容待人工审查；保留既有章节。")
        elif args.command == "section-recover":
            result = recover_section(args.project_dir, args.section,
                                     confirm=args.confirm, unlock_stale=args.unlock_stale)
            print(f"章节 {result['section_id']} 已回退至本次 AI 写入前状态，可重新运行写作。")
        elif args.command == "approve-section":
            result = approve_section(args.project_dir, args.section,
                                     acknowledgement=args.acknowledgement)
            print(f"章节 {result['section_id']}：已由操作者声明审核通过；仍需全篇终审。")
        elif args.command == "outline-init":
            candidate = outline_template(args.project_dir)
            write_new_report(args.output, candidate)
            print(f"层级目录候选已生成：{args.output}。需要人工核实并编辑 roots 后才能应用。")
        elif args.command == "outline-plan":
            report = outline_plan(args.project_dir, args.proposal, apply=args.apply)
            if args.output:
                write_json(args.output, report)
            print(f"目录方案：{report['status']}，一级节点 {report['root_count']} 个，"
                  f"叶子章节 {report['section_count']} 个；"
                  f"已锁定：{report['locked']}")
        elif args.command == "outline-lock":
            result = lock_outline(args.project_dir, confirm=args.confirm)
            print(f"一级目录锁：{result['status']}，指纹：{result['sha256']}")
        elif args.command == "review-init":
            template = review_template(args.suggestions)
            write_new_report(args.output, template)
            multi = len(template.get("ambiguous_decisions", []))
            print(f"审核清单：{args.output}；确定性候选 {len(template['decisions'])} 项"
                  f"，多档评分需逐项审查 {multi} 项。多档项目不得自动采用单一分值。")
        elif args.command == "review-recover":
            result = recover_review(args.project_dir, confirm=args.confirm)
            print(f"已恢复审核导入事务：{', '.join(result['restored_files'])}；"
                  f"保留原始备份 {result['backup']}")
        elif args.command == "review-import":
            plan = reviewed_import(args.project_dir, args.suggestions, args.review,
                                   args.index, args.source_dir, apply=args.apply)
            if args.output:
                write_json(args.output, plan)
            print(f"{'已登记' if args.apply else '仅预览'}：新增 {plan['statistics']['new_requirements']} 条，"
                  f"多档待处理 {plan['statistics'].get('ambiguous_pending', 0)} 条，"
                  f"拒绝 {plan['statistics']['rejected']} 条，"
                  f"待审 {plan['statistics']['pending']} 条，"
                  f"重复 {plan['statistics']['duplicate']} 条。"
                  + (f"备份：{plan.get('backup_dir')}" if plan.get('backup_dir') else ""))
        elif args.command == "analyze-ai":
            result = analyze_tender(args.index, args.source_dir, model=args.model,
                                    endpoint=args.endpoint, allow_remote=args.allow_remote,
                                    api_key_env=args.api_key_env, timeout=args.timeout)
            write_new_report(args.output, result)
            print(f"AI 分析建议 {len(result['candidates'])} 项；"
                  f"选用证据块 {result['statistics']['selected_blocks']} 个；"
                  f"已保存 {args.output}。必须人工复核，不自动修改正式需求。")
        elif args.command == "score-candidates":
            result = discover_score_candidates(args.index, args.source_dir)
            write_new_report(args.output, result)
            print(f"技术评分候选 {len(result['candidates'])} 条；歧义行 "
                  f"{result['statistics']['ambiguous_scoring_lines']} 条。"
                  f"已保存 {args.output}；候选必须人工/Agent 核准。")
        elif args.command in ("check", "audit"):
            report = inspect_with_evidence(args.project_dir, strict=args.command == "audit",
                                           index_paths=args.index, source_dir=args.source_dir,
                                           require_evidence=args.command == "audit")
            if args.output:
                write_json(args.output, report)
            print(f"{'通过' if report['ok'] else '未通过'}｜{report['statistics']}")
            for label, key in (("错误", "errors"), ("提示", "warnings")):
                for entry in report[key]:
                    print(f"{label}：{entry}")
            return 0 if report["ok"] else 1
        elif args.command == "allocate":
            result = allocate(args.project_dir, args.total)
            for sid, count in result.items():
                print(f"{sid}: {count} 字")
        elif args.command == "scaffold":
            result = scaffold(args.project_dir)
            print(f"新建 {len(result)} 个章节文件（不会覆盖已有正文）")
        elif args.command == "export":
            report = export_docx(args.project_dir, args.output, args.draft,
                                 index_paths=args.index, source_dir=args.source_dir)
            print(f"已导出：{args.output}；机械审查：{'通过' if report['ok'] else '未通过'}；"
                  f"证据状态：{report['evidence']['status']}")
        return 0
    except (OSError, ValueError, RuntimeError, KeyError, TypeError) as exc:
        print(f"操作失败：{exc}", file=sys.stderr)
        return 2
