"""Integration release blockers: adversarial source, privacy and approval checks."""
import hashlib
import json
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from biaoshu.ai_tender_analysis import _request_model
from biaoshu.consistency import analyze_consistency
from biaoshu.core import initialize, inspect, write_json
from biaoshu.evidence import verify_evidence


class HardeningRegressionTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.project = self.root / "project"
        initialize(self.project)
        self.source_dir = self.root / "sources"
        self.source_dir.mkdir()
        self.source = self.source_dir / "tender.txt"
        self.source.write_text("技术评分要求写出实施方案\n", encoding="utf-8")
        self.index = self.root / "tender.index.json"
        self.index_obj = {
            "schema_version": 1, "source_file": "tender.txt",
            "sha256": hashlib.sha256(self.source.read_bytes()).hexdigest(),
            "blocks": [{"location": "第1行", "text": "技术评分要求写出实施方案"}],
        }
        write_json(self.index, self.index_obj)
        write_json(self.project / "requirements.json", {"requirements": [
            {"id": "R1", "title": "实施方案", "score": 10,
             "evidence": {"file": "tender.txt", "location": "第1行", "quote": "技术评分要求写出实施方案"},
             "response_sections": ["S1"]},
        ]})
        write_json(self.project / "outline.json", {"sections": [
            {"id": "S1", "title": "实施方案", "requirement_ids": ["R1"],
             "weight": 1, "target_words": 0},
        ]})
        self.chapter = self.project / "sections" / "S1.md"
        self.chapter.write_text("# 实施方案\n\n详细介绍实施任务安排及质量复核机制。\n", encoding="utf-8")

    def test_fake_block_with_real_file_digest_does_not_pass(self):
        index = json.loads(self.index.read_text(encoding="utf-8"))
        index["blocks"][0]["text"] = "不存在于真实招标文件中的虚构人员配置"
        write_json(self.index, index)
        reqs = json.loads((self.project / "requirements.json").read_text(encoding="utf-8"))
        reqs["requirements"][0]["evidence"]["quote"] = "不存在于真实招标文件中的虚构人员配置"
        write_json(self.project / "requirements.json", reqs)
        report = verify_evidence(self.project, [self.index], self.source_dir)
        self.assertFalse(report["ok"])
        self.assertIn("index_content_mismatch", {e["code"] for e in report["errors"]})

    def test_stable_index_passes_when_hashed_and_reextracted(self):
        self.assertTrue(verify_evidence(self.project, [self.index], self.source_dir)["ok"])

    def test_docx_column_metadata_cannot_be_cited_as_source_text(self):
        from docx import Document
        from biaoshu.extraction import extract_document
        doc = Document()
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "评分项"
        table.cell(0, 1).text = "实施方案10分"
        tender = self.source_dir / "table.docx"
        doc.save(tender)
        index = self.root / "table.index.json"
        write_json(index, extract_document(tender).index)
        reqs = json.loads((self.project / "requirements.json").read_text(encoding="utf-8"))
        reqs["requirements"][0]["evidence"] = {
            "file": "table.docx", "location": "第1表第1行", "quote": "列2：实施方案10分",
        }
        write_json(self.project / "requirements.json", reqs)
        report = verify_evidence(self.project, [index], self.source_dir)
        self.assertFalse(report["ok"])
        self.assertIn("quote_mismatch", {e["code"] for e in report["errors"]})
        reqs["requirements"][0]["evidence"]["quote"] = "实施方案10分"
        write_json(self.project / "requirements.json", reqs)
        self.assertTrue(verify_evidence(self.project, [index], self.source_dir)["ok"])

    def test_model_adapter_ignores_proxy_environment(self):
        class Response:
            status = 200
            def __enter__(self):
                return self
            def __exit__(self, *args):
                return False
            def read(self, max_size):
                return b'{"choices":[{"message":{"content":"{}"},"finish_reason":"stop"}]}'
        class FakeOpener:
            def open(self, request, timeout):
                self.request = request
                return Response()
        seen = []
        def capture(*handlers):
            seen.extend(handlers)
            return FakeOpener()
        with patch.dict(os.environ, {"HTTP_PROXY": "http://untrusted-proxy.example:8888",
                                     "HTTPS_PROXY": "http://untrusted-proxy.example:8888",
                                     "NO_PROXY": ""}), patch(
                "biaoshu.ai_tender_analysis.urllib.request.build_opener", side_effect=capture):
            result = _request_model(
                "http://127.0.0.1:11434/v1/chat/completions", "local-llm", "confidential tender",
                api_key="fake-local-secret", timeout=3)
        self.assertEqual(result, "{}")
        proxy_handlers = [x for x in seen
                          if isinstance(x, __import__("urllib.request", fromlist=["ProxyHandler"]).ProxyHandler)]
        self.assertEqual(len(proxy_handlers), 1)
        self.assertEqual(proxy_handlers[0].proxies, {})

    def test_missing_checkpoint_never_approves_banner_draft(self):
        self.chapter.write_text(
            "# 实施方案\n\n> AI生成草稿（待人工审查）\n\n以下是详细的实施计划、过程管理与质量复核机制。\n",
            encoding="utf-8",
        )
        report = inspect(self.project, strict=True)
        self.assertFalse(report["ok"])
        self.assertTrue(any("待审核标记" in e for e in report["errors"]))

    def test_bounded_values_do_not_create_false_conflicts(self):
        write_json(self.project / "facts.json", {"facts": [
            {"key": "response_time", "status": "verified", "value": "不超过2小时",
             "evidence": {"file": "tender.txt", "location": "第1行",
                          "quote": "技术评分要求写出实施方案"}},
        ]})
        self.chapter.write_text("# 实施方案\n\n响应时间为2小时内，我们将在时限内响应。\n",
                                encoding="utf-8")
        self.assertTrue(analyze_consistency(self.project)["ok"])
        self.chapter.write_text("# 实施方案\n\n响应时间为1小时，安排人员及时响应。\n",
                                encoding="utf-8")
        self.assertTrue(analyze_consistency(self.project)["ok"])
        self.assertEqual(analyze_consistency(self.project)["statistics"]["errors"], 0)

    def test_calendar_date_not_parsed_as_year_duration(self):
        self.chapter.write_text(
            "# 实施方案\n\n服务期限为2026年1月1日至2027年12月31日。\n", encoding="utf-8")
        self.assertEqual(analyze_consistency(self.project)["statistics"]["commitments_detected"], 0)

    def test_subschedule_is_not_global_project_duration(self):
        self.chapter.write_text(
            "# 实施方案\n\n设备供货工期为15日，项目工期为30日。\n", encoding="utf-8")
        report = analyze_consistency(self.project)
        self.assertTrue(report["ok"])
        self.assertEqual(report["statistics"]["commitments_detected"], 1)


if __name__ == "__main__":
    unittest.main()
