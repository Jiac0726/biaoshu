"""Synthetic end-to-end acceptance tests: every production CLI stage, zero network.

All tender content and model outputs here are fictional. The test simulates
reviewer acknowledgements strictly to exercise gates, not to certify a bid.
"""
import contextlib
import io
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from biaoshu.cli import main
from biaoshu.core import read_json, write_json
from biaoshu.review_import import ACKNOWLEDGEMENT
from biaoshu.section_writer import BANNER, CONFIRM


class SyntheticBidPipelineTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)
        self.project = self.root / "workspace"
        self.sources = self.root / "source"
        self.sources.mkdir()
        self.original = self.sources / "fictional-tender.txt"
        self.original.write_text(
            "技术评分\n"
            "实施组织方案：10分\n"
            "质量保障措施：15分\n"
            "项目工期为30日\n"
            "人员配置另行核实\n",
            encoding="utf-8",
        )
        self.extracted = self.root / "extracted.md"
        self.index = self.extracted.with_suffix(".index.json")
        self.candidates = self.root / "ai-suggestions.json"
        self.scores = self.root / "score-candidates.json"
        self.review = self.root / "review.json"
        self.proposal = self.root / "outline-proposal.json"
        self.outfile = self.root / "official.docx"
        self.model_calls = []
        self.assertEqual(self.cli("init", str(self.project)), 0)
        self.assertEqual(self.cli("extract", str(self.original), "--output", str(self.extracted)), 0)

    @staticmethod
    def cli(*args):
        with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
            return main(list(args))

    def source_args(self):
        return ("--index", str(self.index), "--source-dir", str(self.sources))

    def fake_analysis_model(self, endpoint, model, prompt, *, api_key, timeout):
        self.model_calls.append(("analyze", model))
        self.assertEqual(endpoint, "http://127.0.0.1:11434/v1/chat/completions")
        self.assertIn("实施组织方案：10分", prompt)
        return json.dumps({"items": [
            {"kind": "technical_score", "title": "实施组织方案", "score": 10,
             "evidence": {"file": self.original.name, "location": "第2行",
                          "quote": "实施组织方案：10分"},
             "rationale": "虚构评分表明确列出实施组织"},
            {"kind": "technical_score", "title": "质量保障措施", "score": 15,
             "evidence": {"file": self.original.name, "location": "第3行",
                          "quote": "质量保障措施：15分"},
             "rationale": "虚构评分表明确列出质量措施"},
        ]}, ensure_ascii=False)

    def fake_writing_model(self, endpoint, model, prompt, *, api_key, timeout, instructions):
        self.model_calls.append(("write", model))
        self.assertEqual(endpoint, "http://127.0.0.1:11434/v1/chat/completions")
        self.assertIn("绝不虚构人员", instructions)
        payload = json.loads(prompt.split("\n", 1)[1])
        self.assertEqual(len(payload["requirements"]), 1)
        self.assertNotIn("虚构的未批准人员", prompt)
        self.assertEqual(payload["facts"][0]["value"], "30日")
        rid = payload["requirements"][0]["id"]
        sid = payload["section_id"]
        marker = {"S1": "【S1专属：组织调度与里程碑安排】",
                  "S2": "【S2专属：质量抽检与闭环整改】"}[sid]
        return json.dumps({
            "body": (
                "按照项目评分要求，采用负责人统筹、工作组执行、质量人员复核的工作机制。"
                "项目工期为30日，按照任务节点安排实施检查，并记录问题、整改过程和交付资料。"
                "所有岗位安排以最终审核的人力配置为准，未经确认不写入具体人员名单。"
                + marker
            ),
            "used_requirement_ids": [rid], "used_fact_keys": ["project_duration"],
        }, ensure_ascii=False)

    def prepare_suggestions(self):
        self.assertEqual(self.cli("score-candidates", *self.source_args(),
                                  "--output", str(self.scores)), 0)
        found = read_json(self.scores)
        self.assertEqual([c["title"] for c in found["candidates"]],
                         ["实施组织方案", "质量保障措施"])
        with patch("biaoshu.ai_tender_analysis._request_model",
                   side_effect=self.fake_analysis_model):
            self.assertEqual(self.cli("analyze-ai", *self.source_args(),
                                      "--model", "fictional-local-model",
                                      "--output", str(self.candidates)), 0)
        payload = read_json(self.candidates)
        self.assertEqual(payload["status"], "suggestions_only")
        self.assertEqual([item["status"] for item in payload["candidates"]],
                         ["pending_review", "pending_review"])
        self.assertFalse((self.project / "requirements.json").read_text(encoding="utf-8").find('"R1"') >= 0)

    def prepare_reviewed_outline(self):
        self.prepare_suggestions()
        self.assertEqual(self.cli("review-init", "--suggestions", str(self.candidates),
                                  "--output", str(self.review)), 0)
        review = read_json(self.review)
        self.assertEqual([x["decision"] for x in review["decisions"]], ["pending", "pending"])
        immutable_before = {
            name: (self.project / name).read_bytes()
            for name in ("project.json", "requirements.json", "outline.json", "facts.json")
        }
        self.assertEqual(self.cli("review-import", str(self.project),
                                  "--suggestions", str(self.candidates),
                                  "--review", str(self.review),
                                  *self.source_args(), "--apply"), 2)
        for name, original in immutable_before.items():
            self.assertEqual((self.project / name).read_bytes(), original,
                             f"rejected pending review mutated {name}")
        self.assertFalse((self.project / "reports" / "review-import-transaction.json").exists())
        self.assertFalse(list((self.project / "reports").glob("review-backups/*")))
        for decision in review["decisions"]:
            decision["decision"] = "accept"
        review["acknowledgement"] = ACKNOWLEDGEMENT
        write_json(self.review, review)

        original_requirements = (self.project / "requirements.json").read_bytes()
        self.assertEqual(self.cli("review-import", str(self.project),
                                  "--suggestions", str(self.candidates),
                                  "--review", str(self.review),
                                  *self.source_args()), 0)
        self.assertEqual((self.project / "requirements.json").read_bytes(), original_requirements)
        self.assertEqual(self.cli("review-import", str(self.project),
                                  "--suggestions", str(self.candidates),
                                  "--review", str(self.review),
                                  *self.source_args(), "--apply"), 0)
        requirements = read_json(self.project / "requirements.json")["requirements"]
        self.assertEqual([x["id"] for x in requirements], ["R1", "R2"])
        self.assertEqual([x["response_sections"] for x in requirements], [["S1"], ["S2"]])
        write_json(self.project / "facts.json", {"facts": [
            {"key": "project_duration", "status": "verified", "value": "30日",
             "evidence": {"file": self.original.name, "location": "第4行",
                          "quote": "项目工期为30日"}},
            {"key": "staff", "status": "pending", "value": "虚构的未批准人员"},
        ]})

        self.assertEqual(self.cli("outline-init", str(self.project),
                                  "--output", str(self.proposal)), 0)
        proposal = read_json(self.proposal)
        proposal["roots"] = [{"id": "G1", "title": "项目实施与质量保障",
                              "children": [{"section_id": "S1"}, {"section_id": "S2"}]}]
        write_json(self.proposal, proposal)
        self.assertEqual(self.cli("outline-plan", str(self.project), "--proposal",
                                  str(self.proposal)), 0)
        self.assertEqual(self.cli("outline-plan", str(self.project), "--proposal",
                                  str(self.proposal), "--apply"), 0)
        self.assertEqual(self.cli("outline-lock", str(self.project), "--confirm"), 0)
        self.assertEqual(self.cli("allocate", str(self.project), "--total", "100"), 0)
        self.assertEqual(self.cli("scaffold", str(self.project)), 0)

    def write_all(self, writer=None):
        with patch("biaoshu.section_writer._request_model",
                   side_effect=writer or self.fake_writing_model):
            return self.cli("write-sections", str(self.project), *self.source_args(),
                            "--model", "fictional-local-model")

    def approve_all(self):
        for sid in ("S1", "S2"):
            self.assertEqual(self.cli("approve-section", str(self.project),
                                      "--section", sid,
                                      "--acknowledgement", CONFIRM), 0)

    def test_entire_pipeline_from_original_to_reviewed_docx(self):
        from docx import Document
        self.prepare_reviewed_outline()
        self.assertEqual(self.write_all(), 0)
        self.assertIn(BANNER, (self.project / "sections" / "S1.md").read_text(encoding="utf-8"))
        audit = self.root / "before-audit.json"
        self.assertEqual(self.cli("audit", str(self.project), *self.source_args(),
                                  "--output", str(audit)), 1)
        self.assertEqual(read_json(audit)["evidence"]["status"], "verified")
        self.assertFalse(read_json(audit)["ok"])
        self.assertNotEqual(self.cli("export", str(self.project), *self.source_args(),
                                     "--output", str(self.outfile)), 0)
        self.assertFalse(self.outfile.exists())
        self.approve_all()
        self.assertEqual(self.cli("consistency-audit", str(self.project)), 0)
        checked = self.root / "check-clean.json"
        self.assertEqual(self.cli("check", str(self.project), "--output", str(checked)), 0)
        self.assertTrue(read_json(checked)["ok"])
        final_audit = self.root / "final-audit.json"
        self.assertEqual(self.cli("audit", str(self.project), *self.source_args(),
                                  "--output", str(final_audit)), 0)
        result = read_json(final_audit)
        self.assertTrue(result["ok"])
        self.assertEqual(result["evidence"]["status"], "verified")
        self.assertEqual(result["evidence"]["summary"]["matched"], 3)
        self.assertEqual(result["consistency"]["statistics"]["errors"], 0)
        self.assertEqual(self.cli("export", str(self.project), *self.source_args(),
                                  "--output", str(self.outfile)), 0)
        doc = Document(str(self.outfile))
        headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
        self.assertEqual(headings[:3], ["项目实施与质量保障", "实施组织方案", "质量保障措施"])
        full_text = "\n".join(p.text for p in doc.paragraphs)
        self.assertIn("项目工期为30日", full_text)
        self.assertNotIn(BANNER, full_text)
        markers = {"实施组织方案": "【S1专属：组织调度与里程碑安排】",
                   "质量保障措施": "【S2专属：质量抽检与闭环整改】"}
        for title, marker in markers.items():
            heading_index = next(i for i, p in enumerate(doc.paragraphs)
                                 if p.text == title and p.style.name.startswith("Heading"))
            next_heading = next((i for i in range(heading_index + 1, len(doc.paragraphs))
                                 if doc.paragraphs[i].style.name.startswith("Heading")),
                                len(doc.paragraphs))
            chapter_text = "\n".join(p.text for p in doc.paragraphs[heading_index+1:next_heading])
            self.assertIn(marker, chapter_text, f"missing {title} content")
            for other_marker in markers.values():
                self.assertEqual(other_marker in chapter_text, other_marker == marker,
                                 f"wrong section body assigned to {title}")
            self.assertEqual(full_text.count(marker), 1, "duplicated or omitted section body")
        self.assertEqual(len([c for c in self.model_calls if c[0] == "write"]), 2)

    def test_e2e_response_trace_review_gate_and_current_paragraph_locations(self):
        from biaoshu.response_trace import ACKNOWLEDGEMENT as TRACE_ACK, trace_path
        self.prepare_reviewed_outline()
        self.assertEqual(self.write_all(), 0)
        self.approve_all()
        self.assertEqual(self.cli("trace-init", str(self.project)), 0)
        self.assertEqual(self.cli("audit", str(self.project), *self.source_args()), 1)
        trace = read_json(trace_path(self.project))
        self.assertEqual(len(trace["items"]), 2)
        for row in trace["items"]:
            body = (self.project / "sections" / (row["section_id"] + ".md")).read_text(encoding="utf-8")
            row["line"] = 3
            row["quote"] = body.splitlines()[2]
            row["decision"] = "confirmed"
        trace["acknowledgement"] = TRACE_ACK
        write_json(trace_path(self.project), trace)
        self.assertEqual(self.cli("trace-audit", str(self.project)), 0)
        self.assertEqual(self.cli("audit", str(self.project), *self.source_args()), 0)
        self.assertEqual(self.cli("export", str(self.project), *self.source_args(),
                                  "--output", str(self.outfile)), 0)
        self.assertTrue(self.outfile.is_file())

    def test_e2e_semantic_review_stays_advisory_and_source_locked(self):
        from biaoshu.response_trace import ACKNOWLEDGEMENT as TRACE_ACK, trace_path
        self.prepare_reviewed_outline()
        self.assertEqual(self.write_all(), 0)
        self.approve_all()
        self.assertEqual(self.cli("trace-init", str(self.project)), 0)
        register = read_json(trace_path(self.project))
        register["acknowledgement"] = TRACE_ACK
        for item in register["items"]:
            row = (self.project / "sections" / (item["section_id"] + ".md")).read_text(encoding="utf-8").splitlines()
            item.update(line=3, quote=row[2], decision="confirmed")
        write_json(trace_path(self.project), register)
        before = trace_path(self.project).read_bytes()
        def semantic(endpoint, model, prompt, *, api_key, timeout, instructions):
            data = json.loads(prompt.split("\n", 1)[1])
            return json.dumps({
                "verdict": "partial", "severity": "medium",
                "reason": "已经提出部分实施措施，但还需由审稿人核查评分细则。",
                "suggestion": "请进一步补充岗位职责、关键节点、过程检查记录与验收依据。",
                "source_quote": data["tender_source"][:18],
                "response_quote": data["response_text"][:18],
            }, ensure_ascii=False)
        output = self.root / "semantic-suggestions.json"
        with patch("biaoshu.semantic_review._request_model", side_effect=semantic):
            self.assertEqual(self.cli(
                "semantic-review", str(self.project), *self.source_args(),
                "--model", "fictional-local-model", "--output", str(output)), 0)
        report = read_json(output)
        self.assertEqual(report["statistics"]["examined_links"], 2)
        self.assertTrue(all(x["status"] == "pending_human_review" for x in report["items"]))
        self.assertEqual(trace_path(self.project).read_bytes(), before)
        self.assertEqual(self.cli("audit", str(self.project), *self.source_args()), 0)

    def test_failure_then_resume_does_not_overwrite_first_chapter(self):
        self.prepare_reviewed_outline()
        def fails_second(endpoint, model, prompt, *, api_key, timeout, instructions):
            payload = json.loads(prompt.split("\n", 1)[1])
            if payload["section_id"] == "S2":
                raise RuntimeError("synthetic network interruption")
            return self.fake_writing_model(endpoint, model, prompt, api_key=api_key,
                                           timeout=timeout, instructions=instructions)
        self.assertEqual(self.write_all(fails_second), 2)
        first = (self.project / "sections" / "S1.md").read_bytes()
        self.assertEqual(read_json(self.project / "reports" / "section-writing" / "S2.json")["status"],
                         "failed")
        self.assertEqual(self.write_all(), 0)
        self.assertEqual((self.project / "sections" / "S1.md").read_bytes(), first)
        self.assertEqual(len([c for c in self.model_calls if c[0] == "write"]), 2)
        self.approve_all()
        self.assertEqual(self.cli("audit", str(self.project), *self.source_args()), 0)

    def test_changed_source_blocks_completed_bid_export(self):
        self.prepare_reviewed_outline()
        self.assertEqual(self.write_all(), 0)
        self.approve_all()
        self.original.write_text("已替换的虚构招标文本", encoding="utf-8")
        self.assertEqual(self.cli("audit", str(self.project), *self.source_args()), 1)
        self.assertNotEqual(self.cli("export", str(self.project), *self.source_args(),
                                     "--output", str(self.outfile)), 0)
        self.assertFalse(self.outfile.exists())

    def test_consistency_blocks_without_any_ai_approval_sidecar(self):
        """Prove audit/export rejection is caused by consistency, not an obsolete AI SHA."""
        self.prepare_reviewed_outline()
        one = self.project / "sections" / "S1.md"
        two = self.project / "sections" / "S2.md"
        one.write_text("# 实施组织方案\n\n项目工期为30日，按工作节点推进项目实施与过程检查。\n",
                       encoding="utf-8")
        two.write_text("# 质量保障措施\n\n项目工期为30日，按质量标准落实资料复核和成果审定。\n",
                       encoding="utf-8")
        self.assertEqual(self.cli("audit", str(self.project), *self.source_args()), 0)
        two.write_text("# 质量保障措施\n\n项目工期为45日，按质量标准落实资料复核和成果审定。\n",
                       encoding="utf-8")
        for sid in ("S1", "S2"):
            self.assertFalse((self.project / "reports" / "section-writing"
                              / (sid + ".json")).exists())
        normal = self.root / "check-conflict.json"
        self.assertEqual(self.cli("check", str(self.project), "--output", str(normal)), 0)
        check_result = read_json(normal)
        self.assertTrue(check_result["ok"])  # normal check warns rather than hard-rejects
        self.assertTrue(any("cross_section_conflict" in w
                            for w in check_result["warnings"]))
        conflict_audit = self.root / "audit-conflict.json"
        self.assertEqual(self.cli("audit", str(self.project), *self.source_args(),
                                  "--output", str(conflict_audit)), 1)
        strict = read_json(conflict_audit)
        self.assertFalse(strict["ok"])
        self.assertGreater(strict["consistency"]["statistics"]["errors"], 0)
        self.assertTrue(any(x["code"] == "cross_section_conflict"
                            for x in strict["consistency"]["findings"]))
        self.assertTrue(any("cross_section_conflict" in e for e in strict["errors"]))
        self.assertFalse(any("审批" in e or "正文已更改" in e for e in strict["errors"]))
        self.assertNotEqual(self.cli("export", str(self.project), *self.source_args(),
                                     "--output", str(self.outfile)), 0)
        self.assertFalse(self.outfile.exists())

    def test_unreviewed_edit_and_conflicting_commitment_are_rejected(self):
        self.prepare_reviewed_outline()
        self.assertEqual(self.write_all(), 0)
        self.approve_all()
        section = self.project / "sections" / "S2.md"
        section.write_text(section.read_text(encoding="utf-8") +
                           "\n项目工期为45日，实施分阶段交付。\n", encoding="utf-8")
        conflicts = self.root / "conflicts.json"
        self.assertEqual(self.cli("consistency-audit", str(self.project),
                                  "--output", str(conflicts)), 1)
        self.assertGreater(read_json(conflicts)["statistics"]["errors"], 0)
        self.assertEqual(self.cli("audit", str(self.project), *self.source_args()), 1)
        self.assertNotEqual(self.cli("export", str(self.project), *self.source_args(),
                                     "--output", str(self.outfile)), 0)
        self.assertFalse(self.outfile.exists())


if __name__ == "__main__":
    unittest.main()
