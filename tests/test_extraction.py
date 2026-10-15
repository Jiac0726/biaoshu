"""Regression tests for source order, page/row locators and index integrity."""
import hashlib
import json
import tempfile
import unittest
from pathlib import Path

from biaoshu.extraction import extract_document


class ExtractionTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)

    def tearDown(self):
        self.temp.cleanup()

    def test_text_preserves_original_line_numbers_and_hash(self):
        source = self.root / "招标要求.txt"
        source.write_text("第一行\n\n技术参数：A|B\n", encoding="utf-8")
        result = extract_document(source)
        self.assertEqual([b["location"] for b in result.index["blocks"]], ["第1行", "第3行"])
        self.assertEqual(result.index["blocks"][1]["text"], "技术参数：A|B")
        self.assertEqual(result.index["sha256"], hashlib.sha256(source.read_bytes()).hexdigest())
        self.assertIn("技术参数：A|B", result.markdown)

    def test_empty_file_does_not_produce_false_evidence(self):
        source = self.root / "empty.txt"
        source.write_text("\n", encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "没有可提取文本"):
            extract_document(source)

    def test_missing_doc_dependencies_show_valid_pip_install_command(self):
        from unittest.mock import patch
        import sys
        for suffix, module in ((".pdf", "pypdf"), (".docx", "docx")):
            with self.subTest(suffix=suffix):
                source = self.root / ("missing" + suffix)
                source.write_bytes(b"not a real source because dependency is mocked")
                with patch.dict(sys.modules, {module: None}):
                    with self.assertRaisesRegex(RuntimeError, r'pip install -e "\\.\\[docs\\]"'):
                        extract_document(source)

    def test_docx_keeps_paragraph_table_paragraph_order(self):
        try:
            from docx import Document
        except ImportError:
            self.skipTest("python-docx optional dependency unavailable")
        source = self.root / "test.docx"
        doc = Document()
        doc.add_paragraph("第一段：技术范围")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "要求"
        table.cell(0, 1).text = "指标|阈值"
        table.cell(1, 0).text = "工期"
        table.cell(1, 1).text = "30日"
        doc.add_paragraph("第二段：交付标准")
        doc.save(source)
        result = extract_document(source)
        self.assertEqual([b["location"] for b in result.index["blocks"]],
                         ["第1段", "第1表第1行", "第1表第2行", "第2段"])
        self.assertIn("指标\\|阈值", result.markdown)
        self.assertLess(result.markdown.index("第一段：技术范围"), result.markdown.index("工期"))
        self.assertLess(result.markdown.index("工期"), result.markdown.index("第二段：交付标准"))

    def test_pdf_page_locator_and_blank_page_warning(self):
        try:
            from pypdf import PdfWriter
            from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
        except ImportError:
            self.skipTest("pypdf optional dependency unavailable")
        source = self.root / "test.pdf"
        writer = PdfWriter()
        page = writer.add_blank_page(width=200, height=200)
        font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                                 NameObject("/Subtype"): NameObject("/Type1"),
                                 NameObject("/BaseFont"): NameObject("/Helvetica")})
        font_ref = writer._add_object(font)
        page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"):
                                                             DictionaryObject({NameObject("/F1"): font_ref})})
        stream = DecodedStreamObject()
        stream.set_data(b"BT /F1 12 Tf 20 100 Td (Bid technical requirements) Tj ET")
        page[NameObject("/Contents")] = writer._add_object(stream)
        writer.add_blank_page(width=200, height=200)
        with source.open("wb") as handle:
            writer.write(handle)
        result = extract_document(source)
        self.assertEqual([b["location"] for b in result.index["blocks"]], ["第1页"])
        self.assertIn("Bid technical requirements", result.index["blocks"][0]["text"])
        self.assertTrue(any("第2页" in w for w in result.index["warnings"]))
        self.assertIn("## 第2页", result.markdown)

    def test_cli_writes_companion_index(self):
        from biaoshu.cli import main
        source = self.root / "tender.txt"
        output = self.root / "tender-extracted.md"
        source.write_text("技术评分：实施组织\n", encoding="utf-8")
        self.assertEqual(main(["extract", str(source), "--output", str(output)]), 0)
        index = json.loads(output.with_suffix(".index.json").read_text(encoding="utf-8"))
        self.assertEqual(index["source_file"], "tender.txt")
        self.assertEqual(index["blocks"][0]["location"], "第1行")
        self.assertIn("技术评分：实施组织", output.read_text(encoding="utf-8"))


if __name__ == "__main__":
    unittest.main()
