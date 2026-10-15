"""Real TXT/source-index integration tests for the primary CLI evidence gate."""
import contextlib
import io
import json
import tempfile
import unittest
from pathlib import Path

from biaoshu.cli import export_docx, main
from biaoshu.core import initialize, write_json


class CliEvidenceIntegrationTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.sources = self.root / "sources"
        self.sources.mkdir()
        self.project = self.root / "project"
        initialize(self.project)
        self.source = self.sources / "tender.txt"
        self.source.write_text("技术评分：实施组织\n交付时间：30日\n", encoding="utf-8")
        self.extracted = self.root / "tender-extracted.md"
        self.assertEqual(self.run_cli("extract", str(self.source), "--output", str(self.extracted)), 0)
        self.index = self.extracted.with_suffix(".index.json")
        write_json(self.project / "requirements.json", {"requirements": [
            {"id": "R1", "title": "实施组织", "score": 10,
             "evidence": {"file": "tender.txt", "location": "第1行", "quote": "技术评分：实施组织"},
             "response_sections": ["S1"]},
        ]})
        write_json(self.project / "facts.json", {"facts": [
            {"key": "delivery_time", "value": "30日", "status": "verified",
             "evidence": {"file": "tender.txt", "location": "第2行", "quote": "交付时间：30日"}}
        ]})
        write_json(self.project / "outline.json", {"sections": [
            {"id": "S1", "title": "实施组织", "requirement_ids": ["R1"],
             "weight": 1.0, "target_words": 0}
        ]})
        (self.project / "sections" / "S1.md").write_text(
            "# 实施组织\n根据招标要求设置实施组织和项目管理机制。\n", encoding="utf-8"
        )

    @staticmethod
    def run_cli(*args):
        with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
            return main(list(args))

    def evidence_args(self):
        return ("--index", str(self.index), "--source-dir", str(self.sources))

    def test_audit_requires_evidence(self):
        self.assertEqual(self.run_cli("audit", str(self.project)), 1)

    def test_check_remains_compatible_but_reports_unchecked(self):
        report = self.root / "check.json"
        self.assertEqual(self.run_cli("check", str(self.project), "--output", str(report)), 0)
        saved = json.loads(report.read_text(encoding="utf-8"))
        self.assertEqual(saved["evidence"]["status"], "not_checked")
        self.assertTrue(any("未核验" in warning for warning in saved["warnings"]))

    def test_valid_audit_records_provenance(self):
        report = self.root / "audit.json"
        self.assertEqual(self.run_cli("audit", str(self.project), *self.evidence_args(),
                                      "--output", str(report)), 0)
        saved = json.loads(report.read_text(encoding="utf-8"))
        self.assertEqual(saved["evidence"]["status"], "verified")
        self.assertEqual(saved["evidence"]["summary"]["matched"], 2)

    def test_incomplete_configuration_fails(self):
        self.assertEqual(self.run_cli("check", str(self.project),
                                      "--index", str(self.index)), 1)

    def test_changed_source_fails_audit_and_prevents_official_export(self):
        self.source.write_text("技术评分：变更内容\n交付时间：90日\n", encoding="utf-8")
        out = self.root / "official.docx"
        self.assertEqual(self.run_cli("audit", str(self.project), *self.evidence_args()), 1)
        self.assertNotEqual(self.run_cli("export", str(self.project), *self.evidence_args(),
                                         "--output", str(out)), 0)
        self.assertFalse(out.exists())

    def test_missing_provenance_blocks_official_export(self):
        out = self.root / "no-index.docx"
        self.assertNotEqual(self.run_cli("export", str(self.project), "--output", str(out)), 0)
        self.assertFalse(out.exists())
        with self.assertRaises(ValueError):
            export_docx(self.project, out, draft=False)
        self.assertFalse(out.exists())

    def test_verified_evidence_allows_official_word_export(self):
        try:
            import docx  # noqa: F401
        except ImportError:
            self.skipTest("python-docx optional dependency unavailable")
        out = self.root / "official.docx"
        self.assertEqual(self.run_cli("export", str(self.project), *self.evidence_args(),
                                      "--output", str(out)), 0)
        self.assertTrue(out.is_file())
        from docx import Document
        self.assertIn("实施组织", "\n".join(p.text for p in Document(str(out)).paragraphs))

    def test_draft_works_without_evidence_and_is_visibly_marked(self):
        try:
            from docx import Document
        except ImportError:
            self.skipTest("python-docx optional dependency unavailable")
        out = self.root / "draft.docx"
        self.assertEqual(self.run_cli("export", str(self.project), "--draft",
                                      "--output", str(out)), 0)
        document = Document(str(out))
        self.assertIn("草稿，未完成校验", document.paragraphs[0].text)

    def test_draft_rejects_unsafe_section_identifier(self):
        contents = json.loads((self.project / "outline.json").read_text(encoding="utf-8"))
        contents["sections"][0]["id"] = "../secret"
        write_json(self.project / "outline.json", contents)
        out = self.root / "unsafe.docx"
        self.assertNotEqual(self.run_cli("export", str(self.project), "--draft",
                                         "--output", str(out)), 0)
        self.assertFalse(out.exists())


if __name__ == "__main__":
    unittest.main()
