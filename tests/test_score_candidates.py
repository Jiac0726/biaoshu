"""Source-locked technical scoring proposal tests (no automatic acceptance)."""
import hashlib
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from biaoshu.score_candidates import discover_score_candidates


class ScoreCandidatesTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.src = self.root / "sources"
        self.src.mkdir()
        self.file = self.src / "tender.txt"
        self.file.write_text("技术评分\n实施组织方案：10分\n质量保障措施：15分\n价格评分：30分\n交付期限：30日\n", encoding="utf-8")
        self.index = self.root / "tender.index.json"
        self.blocks = [
            {"location": f"第{i}行", "kind": "line", "text": line}
            for i, line in enumerate(self.file.read_text(encoding="utf-8").splitlines(), 1)
        ]
        self.save_index()

    def save_index(self):
        self.index.write_text(json.dumps({
            "schema_version": 1, "source_file": self.file.name,
            "sha256": hashlib.sha256(self.file.read_bytes()).hexdigest(),
            "blocks": self.blocks, "warnings": [],
        }, ensure_ascii=False), encoding="utf-8")

    def run_discovery(self):
        return discover_score_candidates([self.index], self.src)

    def parser_only(self):
        # Synthetic PDF/table formats here test parsing, not provenance.
        # Production source verification remains enabled for end-to-end calls.
        with patch("biaoshu.score_candidates._source_index",
                   return_value=json.loads(self.index.read_text(encoding="utf-8"))):
            return self.run_discovery()

    def test_only_technical_candidates_with_original_quote(self):
        result = self.run_discovery()
        self.assertEqual(result["status"], "suggestions_only")
        self.assertEqual([c["title"] for c in result["candidates"]], ["实施组织方案", "质量保障措施"])
        self.assertEqual([c["score"] for c in result["candidates"]], [10.0, 15.0])
        self.assertEqual(result["candidates"][0]["evidence"],
                         {"file": "tender.txt", "location": "第2行", "quote": "实施组织方案：10分"})
        self.assertTrue(all(c["status"] == "pending_review" for c in result["candidates"]))

    def test_source_changed_is_rejected(self):
        self.file.write_text("变更后的原件", encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "已变化"):
            self.run_discovery()

    def test_forged_scoring_block_with_correct_source_hash_is_rejected(self):
        self.blocks[1]["text"] = "模型伪造的评分表条目：100分"
        self.save_index()
        with self.assertRaisesRegex(ValueError, "索引与招标原件解析结果不一致"):
            self.run_discovery()

    def test_duplicate_indexes_are_rejected(self):
        with self.assertRaisesRegex(ValueError, "同名原件"):
            discover_score_candidates([self.index, self.index], self.src)

    def test_multiline_docx_table_tiers_stay_one_ambiguous_review_item(self):
        from docx import Document
        from biaoshu.extraction import extract_document
        table_doc = self.src / "grade-tiers.docx"
        doc = Document()
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "质量保障措施"
        table.cell(0, 1).text = "合格得5分\n优秀得10分"
        doc.save(table_doc)
        index = self.root / "grade.index.json"
        parsed = extract_document(table_doc).index
        index.write_text(json.dumps(parsed, ensure_ascii=False), encoding="utf-8")
        result = discover_score_candidates([index], self.src)
        self.assertEqual(result["candidates"], [])
        self.assertEqual(len(result["ambiguous_candidates"]), 1)
        item = result["ambiguous_candidates"][0]
        self.assertEqual(item["score_options"], [5.0, 10.0])
        self.assertEqual(item["evidence"]["location"], "第1表第1行")
        self.assertIn("优秀得10分", item["evidence"]["quote"])

    def test_table_specific_criterion_not_generic_score_category(self):
        self.blocks = [{"location": "第1表第2行", "kind": "table_row",
                        "text": "列1：技术评分 | 列2：实施组织方案 | 列3：10分"}]
        self.save_index()
        result = self.parser_only()
        self.assertEqual([x["title"] for x in result["candidates"]], ["实施组织方案"])

    def test_per_item_penalties_and_bonuses_require_review(self):
        self.blocks = [
            {"location": "第2行", "kind": "line", "text": "安全措施每缺一项扣2分"},
            {"location": "第3行", "kind": "line", "text": "创新方案每提供一项加2分"},
        ]
        self.save_index()
        result = self.parser_only()
        self.assertFalse(result["candidates"])
        self.assertEqual([x["reason"] for x in result["ambiguous_candidates"]],
                         ["adjustment_rule", "adjustment_rule"])
        self.assertEqual([x["score_options"] for x in result["ambiguous_candidates"]], [[2.0], [2.0]])
        self.assertEqual([x["status"] for x in result["ambiguous_candidates"]],
                         ["requires_manual_scoring_review"] * 2)

    def test_docx_row_base_score_and_penalty_is_adjustment_not_tier(self):
        self.blocks = [{"location": "第1表第3行", "kind": "table_row",
                        "text": "列1：安全措施 | 列2：满分10分，每缺一项扣2分"}]
        self.save_index()
        result = self.parser_only()
        self.assertEqual(result["candidates"], [])
        self.assertEqual(len(result["ambiguous_candidates"]), 1)
        self.assertEqual(result["ambiguous_candidates"][0]["reason"], "adjustment_rule")
        self.assertEqual(result["ambiguous_candidates"][0]["score_options"], [10.0, 2.0])

    def test_above_100_scoring_tier_is_preserved_for_review(self):
        self.blocks = [{"location": "第1行", "kind": "line",
                        "text": "质量保障措施：合格90分，优秀105分"}]
        self.save_index()
        result = self.parser_only()
        self.assertFalse(result["candidates"])
        self.assertEqual(result["ambiguous_candidates"][0]["score_options"], [90.0, 105.0])

    def test_normal_single_criterion_remains_direct_suggestion(self):
        self.assertEqual([x["score"] for x in self.run_discovery()["candidates"]],
                         [10.0, 15.0])

    def test_table_row_preserves_location(self):
        self.blocks = [{"location": "第2表第4行", "kind": "table_row",
                        "text": "列1：实施组织方案 | 列2：10分 | 列3：评审实施方案内容"}]
        self.save_index()
        result = self.parser_only()
        self.assertEqual(len(result["candidates"]), 1)
        self.assertEqual(result["candidates"][0]["title"], "实施组织方案")
        self.assertEqual(result["candidates"][0]["evidence"]["location"], "第2表第4行")

    def test_multi_tier_score_is_ambiguous(self):
        self.blocks = [{"location": "第3页", "kind": "page",
                        "text": "质量保障措施：合格得5分，优秀得10分"}]
        self.save_index()
        result = self.parser_only()
        self.assertFalse(result["candidates"])
        self.assertEqual(result["statistics"]["ambiguous_scoring_lines"], 1)

    def test_multi_tier_values_saved_for_human_review_without_numeric_approval(self):
        self.blocks = [{"location": "第3页", "kind": "page",
                        "text": "质量保障措施：合格得5分，优秀得10分"}]
        self.save_index()
        report = self.parser_only()
        self.assertEqual(report["candidates"], [])
        self.assertEqual(report["statistics"]["multi_tier_lines"], 1)
        item = report["ambiguous_candidates"][0]
        self.assertEqual(item["score_options"], [5.0, 10.0])
        self.assertEqual(item["status"], "requires_manual_scoring_review")
        self.assertEqual(item["evidence"]["location"], "第3页")
        self.assertNotIn("score", item)

    def test_bidding_price_multi_tier_is_not_technical_candidate(self):
        self.blocks = [{"location": "第1页", "kind": "page",
                        "text": "商务评分：报价合格得5分，报价优秀得10分"}]
        self.save_index()
        report = self.parser_only()
        self.assertFalse(report["ambiguous_candidates"])
        self.assertFalse(report["candidates"])

    def test_pdf_page_lines_are_independent(self):
        self.blocks = [{"location": "第7页", "kind": "page",
                        "text": "技术评分\n实施组织方案：10分\n质量保障措施：15分"}]
        self.save_index()
        result = self.parser_only()
        self.assertEqual(len(result["candidates"]), 2)
        self.assertEqual({c["evidence"]["location"] for c in result["candidates"]}, {"第7页"})

    def test_reject_path_traversal_in_index_filename(self):
        data = json.loads(self.index.read_text(encoding="utf-8"))
        data["source_file"] = "../tender.txt"
        self.index.write_text(json.dumps(data), encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "文件名"):
            self.run_discovery()

    def test_no_candidates_does_not_mean_compliant(self):
        self.blocks = [{"location": "第1行", "kind": "line", "text": "投标截止日期"}]
        self.save_index()
        result = self.parser_only()
        self.assertEqual(result["candidates"], [])
        self.assertIn("可能漏检", result["warnings"][0])


if __name__ == "__main__":
    unittest.main()
